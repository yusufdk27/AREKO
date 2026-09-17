import io
from typing import Dict, Any, List, Union, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.detector import sort_resume_chronological




def generate_form_excel(
    output_target: Optional[Union[str, io.BytesIO]] = None,
    header_info: Optional[Dict[str, str]] = None,
    resume_list: Optional[List[Dict[str, Any]]] = None,
    note_oh: str = "",
    nama_so: str = "",
    nama_oh: str = "",
) -> Union[str, io.BytesIO]:
    """
    Membuat file Excel laporan resmi Form Validasi Mutasi Rekening.
    Mendukung penulisan ke path file string atau objek io.BytesIO (in-memory).
    """
    header_info = header_info or {}
    resume_list = resume_list or []
    sorted_resume = sort_resume_chronological(resume_list)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Validasi Mutasi"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Arial", size=11, bold=True)
    font_sec = Font(name="Arial", size=10, bold=True)
    font_bold = Font(name="Arial", size=9, bold=True)
    font_norm = Font(name="Arial", size=9)

    fill_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_blue = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    fill_grey = PatternFill(start_color="808080", end_color="808080", fill_type="solid")
    fill_soft_blue = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

    thin = Side(border_style="thin", color="000000")
    box_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 1. Judul & Header Form Kiri
    ws.merge_cells("C1:G1")
    ws["C1"] = "FORM VALIDASI MUTASI REKENING"
    ws["C1"].font = font_title
    ws["C1"].alignment = Alignment(horizontal="center", vertical="center")

    headers = [
        (3, "Cabang", header_info.get("cabang", "")),
        (4, "Nama Cust", header_info.get("nama_cust", "")),
        (6, "Nomor Rekening", header_info.get("no_rekening", "")),
        (7, "Nama Bank", header_info.get("nama_bank", "")),
        (8, "Nama Pemegang Rekening", header_info.get("nama_pemegang_rek", "")),
    ]

    for r, label, val in headers:
        ws.cell(row=r, column=1, value=label).font = font_norm
        ws.cell(row=r, column=3, value=":").font = font_norm
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=7)
        cell = ws.cell(row=r, column=4, value=val)
        cell.font = font_norm
        cell.fill = fill_green
        for c in range(4, 8):
            ws.cell(row=r, column=c).border = box_border

    ws["A10"] = "Resume Mutasi Rekening"
    ws["A10"].font = font_sec

    # 2. Header Tabel Resume
    ws.merge_cells("A11:A12")
    ws["A11"] = "Bulan"
    ws.merge_cells("B11:C11")
    ws["B11"] = "Frekuensi"
    ws["B12"] = "Debet"
    ws["C12"] = "Kredit"
    ws.merge_cells("D11:E11")
    ws["D11"] = "Mutasi"
    ws["D12"] = "Debet"
    ws["E12"] = "Kredit"
    ws.merge_cells("F11:H11")
    ws["F11"] = "Saldo"
    ws["F12"] = "Tertinggi"
    ws["G12"] = "Rata-Rata"
    ws["H12"] = "Terendah"

    for r in range(11, 13):
        for c in range(1, 9):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill_blue
            cell.font = font_bold
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = box_border

    start_r = 13
    n = len(sorted_resume)
    sum_vals = {"f_db": 0, "f_cr": 0, "m_db": 0.0, "m_cr": 0.0, "s_max": 0.0, "s_avg": 0.0, "s_min": 0.0}

    for i, item in enumerate(sorted_resume):
        curr_r = start_r + i
        b_name = item.get("bulan", "-")

        c1 = ws.cell(row=curr_r, column=1, value=b_name)
        c1.alignment = Alignment(horizontal="center")
        c1.font = font_norm
        c1.border = box_border

        f_db = item.get("freq_db", 0)
        f_cr = item.get("freq_cr", 0)
        m_db = item.get("mutasi_db", 0.0)
        m_cr = item.get("mutasi_cr", 0.0)
        s_max = item.get("saldo_max", 0.0)
        s_avg = item.get("saldo_avg", 0.0)
        s_min = item.get("saldo_min", 0.0)

        sum_vals["f_db"] += f_db
        sum_vals["f_cr"] += f_cr
        sum_vals["m_db"] += m_db
        sum_vals["m_cr"] += m_cr
        sum_vals["s_max"] += s_max
        sum_vals["s_avg"] += s_avg
        sum_vals["s_min"] += s_min

        ws.cell(row=curr_r, column=2, value=f_db).number_format = "#,##0"
        ws.cell(row=curr_r, column=3, value=f_cr).number_format = "#,##0"
        ws.cell(row=curr_r, column=4, value=m_db).number_format = "#,##0.00"
        ws.cell(row=curr_r, column=5, value=m_cr).number_format = "#,##0.00"
        ws.cell(row=curr_r, column=6, value=s_max).number_format = "#,##0.00"
        ws.cell(row=curr_r, column=7, value=s_avg).number_format = "#,##0.00"
        ws.cell(row=curr_r, column=8, value=s_min).number_format = "#,##0.00"

        for c_idx in range(2, 9):
            cell = ws.cell(row=curr_r, column=c_idx)
            cell.font = font_norm
            cell.border = box_border
            cell.alignment = Alignment(horizontal="right")

    end_data_r = start_r + max(len(sorted_resume) - 1, 0)
    avg_r = (end_data_r + 1) if n > 0 else start_r

    # Baris Rata-Rata
    ws.cell(row=avg_r, column=1, value="Rata-Rata")

    avg_f_db = round(sum_vals["f_db"] / n) if n else 0
    avg_f_cr = round(sum_vals["f_cr"] / n) if n else 0
    avg_m_db = (sum_vals["m_db"] / n) if n else 0.0
    avg_m_cr = (sum_vals["m_cr"] / n) if n else 0.0
    avg_s_max = (sum_vals["s_max"] / n) if n else 0.0
    avg_s_avg = (sum_vals["s_avg"] / n) if n else 0.0
    avg_s_min = (sum_vals["s_min"] / n) if n else 0.0

    ws.cell(row=avg_r, column=2, value=avg_f_db).number_format = "#,##0"
    ws.cell(row=avg_r, column=3, value=avg_f_cr).number_format = "#,##0"
    ws.cell(row=avg_r, column=4, value=avg_m_db).number_format = "#,##0.00"
    ws.cell(row=avg_r, column=5, value=avg_m_cr).number_format = "#,##0.00"
    ws.cell(row=avg_r, column=6, value=avg_s_max).number_format = "#,##0.00"
    ws.cell(row=avg_r, column=7, value=avg_s_avg).number_format = "#,##0.00"
    ws.cell(row=avg_r, column=8, value=avg_s_min).number_format = "#,##0.00"

    for c in range(1, 9):
        cell = ws.cell(row=avg_r, column=c)
        cell.fill = fill_green
        cell.font = font_bold
        cell.border = box_border
        if c == 1:
            cell.alignment = Alignment(horizontal="center")
        else:
            cell.alignment = Alignment(horizontal="right")

    # Note OH
    note_label_r = avg_r + 2
    ws.cell(row=note_label_r, column=1, value="Note OH:").font = font_sec
    ws.merge_cells(start_row=note_label_r + 1, start_column=1, end_row=note_label_r + 2, end_column=8)
    note_c = ws.cell(row=note_label_r + 1, column=1, value=note_oh if note_oh else "-")
    note_c.font = font_norm
    note_c.alignment = Alignment(vertical="top")
    for r in range(note_label_r + 1, note_label_r + 3):
        for c in range(1, 9):
            ws.cell(row=r, column=c).border = box_border

    # Tanda Tangan
    sign_r = note_label_r + 4
    ws.cell(row=sign_r, column=1, value="Mengajukan,").font = font_norm
    ws.cell(row=sign_r, column=6, value="Menyetujui,").font = font_norm
    ws.cell(row=sign_r + 4, column=1, value=f"SO: {nama_so}" if nama_so else "SO:").font = font_norm
    ws.cell(row=sign_r + 4, column=6, value=f"Operation Head: {nama_oh}" if nama_oh else "Operation Head:").font = font_norm

    # Set Lebar Kolom Pemisah Awal I, J, K (Kolom 9, 10, 11)
    for c_spacer in [9, 10, 11]:
        ws.column_dimensions[get_column_letter(c_spacer)].width = 2.5

    # 3. Rincian Mutasi Kanan (Mulai Kolom L = Kolom 12, berjarak 7 kolom per bulan sesuai template)
    start_col = 12
    for idx, item in enumerate(sorted_resume):
        b_name = item.get("bulan", f"Bulan {idx + 1}")
        col_b = start_col + (idx * 7)

        ws.cell(row=1, column=col_b, value="Bulan")
        ws.merge_cells(start_row=1, start_column=col_b + 1, end_row=1, end_column=col_b + 2)
        ws.cell(row=1, column=col_b + 1, value=f"Mutasi {b_name}")
        ws.cell(row=1, column=col_b + 3, value="Saldo")

        ws.cell(row=2, column=col_b + 1, value="Debet")
        ws.cell(row=2, column=col_b + 2, value="Kredit")

        for r in range(1, 3):
            for c in range(col_b, col_b + 4):
                cell = ws.cell(row=r, column=c)
                cell.font = font_bold
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = box_border

        # Baris 3: Total Ringkasan Bulan
        ws.cell(row=3, column=col_b, value=b_name).fill = fill_yellow
        ws.cell(row=3, column=col_b).font = font_bold
        ws.cell(row=3, column=col_b).border = box_border

        col_deb = get_column_letter(col_b + 1)
        col_krd = get_column_letter(col_b + 2)

        records = item.get("tx_records", [])
        max_tx_row = max(len(records) + 3, 40)

        ws.cell(row=3, column=col_b + 1, value=f"=SUM({col_deb}4:{col_deb}{max_tx_row})")
        ws.cell(row=3, column=col_b + 2, value=f"=SUM({col_krd}4:{col_krd}{max_tx_row})")

        op_bal = item.get("opening_bal")
        if op_bal is None and records:
            op_bal = records[0]["saldo"] + records[0]["debet"] - records[0]["kredit"]

        ws.cell(row=3, column=col_b + 3, value=op_bal)

        for c in range(col_b, col_b + 4):
            cell = ws.cell(row=3, column=c)
            cell.fill = fill_yellow
            cell.font = font_bold
            cell.border = box_border
            if c > col_b:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0.00"

        # Baris 4 dst: Transaksi Mutasi Langsung
        for r_idx in range(4, max_tx_row + 1):
            tx_idx = r_idx - 4
            rec = records[tx_idx] if tx_idx < len(records) else None

            # Tanggal
            c_bln = ws.cell(row=r_idx, column=col_b, value=rec["date"] if rec else "")
            c_bln.fill = fill_grey
            c_bln.font = Font(name="Arial", size=8, color="FFFFFF")
            c_bln.border = box_border
            c_bln.alignment = Alignment(horizontal="center")

            # Debet
            c_deb = ws.cell(row=r_idx, column=col_b + 1, value=rec["debet"] if (rec and rec["debet"] > 0) else None)
            c_deb.fill = fill_yellow
            c_deb.font = font_norm
            c_deb.border = box_border
            c_deb.number_format = "#,##0.00"

            # Kredit
            c_krd = ws.cell(row=r_idx, column=col_b + 2, value=rec["kredit"] if (rec and rec["kredit"] > 0) else None)
            c_krd.fill = fill_yellow
            c_krd.font = font_norm
            c_krd.border = box_border
            c_krd.number_format = "#,##0.00"

            # Saldo Riil
            c_sal = ws.cell(row=r_idx, column=col_b + 3, value=rec["saldo"] if rec else None)
            c_sal.fill = fill_soft_blue
            c_sal.font = font_norm
            c_sal.border = box_border
            c_sal.number_format = "#,##0.00"
            c_sal.alignment = Alignment(horizontal="right")

        # 3 Kolom Pemisah Kosong Antar Bulan (width 2.5 per kolom sesuai template)
        for c_sep in [col_b + 4, col_b + 5, col_b + 6]:
            ws.column_dimensions[get_column_letter(c_sep)].width = 2.5

        ws.column_dimensions[get_column_letter(col_b)].width = 12
        ws.column_dimensions[get_column_letter(col_b + 1)].width = 16
        ws.column_dimensions[get_column_letter(col_b + 2)].width = 16
        ws.column_dimensions[get_column_letter(col_b + 3)].width = 18


    # Set Column Widths untuk Tabel Form Kiri
    for col in ["A", "B", "C", "D", "E", "F", "G", "H"]:
        ws.column_dimensions[col].width = 15
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 10

    # Save to path or in-memory BytesIO
    if output_target is None:
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf
    elif isinstance(output_target, io.BytesIO):
        wb.save(output_target)
        output_target.seek(0)
        return output_target
    else:
        wb.save(output_target)
        return output_target
